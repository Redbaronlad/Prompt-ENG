import json
import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


MODEL_NAMES = ["GPT-5.6 Sol Ultra", "Fable 5 High", "Grok 4.6 xhigh"]
GROUPS = [
    ("Before (a)", (4, 5, 6)),
    ("Before (b)", (7, 8, 9)),
    ("Before (c)", (10, 11, 12)),
    ("After (a)", (14, 15, 16)),
    ("After (b)", (17, 18, 19)),
    ("After (c)", (20, 21, 22)),
]


def clean(text):
    text = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_generated(text):
    text = clean(text)
    text = re.split(
        r"(?im)^\s*(?:#{1,6}\s*)?(?:\*\*)?"
        r"(?:answers?|answer\s+key|solutions?)\b[^\n]*$",
        text,
        maxsplit=1,
    )[0]
    marker = re.compile(
        r"(?m)^\s*(?:[-*]\s*)?(?:\*\*)?(?:Question\s+)?(\d{1,2})\s*[.):]"
        r"(?:\*\*)?\s*"
    )
    raw_matches = list(marker.finditer(text))
    matches = []
    expected = 1
    for match in raw_matches:
        if int(match.group(1)) == expected:
            matches.append(match)
            expected += 1
            if expected == 11:
                break
    questions = []
    for idx, match in enumerate(matches):
        number = int(match.group(1))
        if not 1 <= number <= 30:
            continue
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
        block = clean(text[match.end():end])
        block = re.sub(r"(?m)^\s*(?:Answer|Rationale|Explanation)\s*:.*$", "", block).strip()
        if block:
            questions.append(block)
    if len(questions) >= 3:
        return questions

    blocks = [clean(x) for x in re.split(r"\n\s*\n", text) if clean(x)]
    likely = [x for x in blocks if "?" in x or re.search(r"\b(?:calculate|solve|identify|determine|explain|analyze|compare|evaluate)\b", x, re.I)]
    return likely or blocks


def split_teacher(text):
    return [clean(x) for x in re.split(r"\n\s*\n", clean(text)) if clean(x)]


def main():
    workbook = Path(os.environ["AP_RESEARCH_WORKBOOK"])
    wb = load_workbook(workbook, read_only=True, data_only=False)
    ws = wb[wb.sheetnames[0]]

    references = {}
    for row in range(20, 24):
        subject = clean(ws.cell(row, 42).value)
        questions = split_teacher(ws.cell(row, 43).value)
        references[subject] = questions

    rows = []
    counts = []
    for row in range(2, 18):
        subject = clean(ws.cell(row, 1).value)
        participant = ws.cell(row, 2).value
        group_payload = []
        for group_name, columns in GROUPS:
            models = []
            for model_name, column in zip(MODEL_NAMES, columns):
                questions = split_generated(ws.cell(row, column).value)
                models.append({"model": model_name, "questions": questions})
                counts.append({
                    "row": row,
                    "subject": subject,
                    "group": group_name,
                    "model": model_name,
                    "count": len(questions),
                })
            group_payload.append({"name": group_name, "models": models})
        rows.append({
            "sheetRow": row,
            "subject": subject,
            "participant": participant,
            "referenceQuestions": references[subject],
            "groups": group_payload,
        })

    payload = {
        "embedding": "Apple NaturalLanguage English 300-dimensional word embedding, mean pooled per question",
        "models": MODEL_NAMES,
        "groups": [name for name, _ in GROUPS],
        "references": references,
        "rows": rows,
        "counts": counts,
    }
    json.dump(payload, sys.stdout, ensure_ascii=False, separators=(",", ":"))


if __name__ == "__main__":
    main()
