"""Reproducible TF-IDF cosine scorer for the workbook's six conditions.

This is the documented reference implementation for lexical cosine. The
original one-off code that populated the live Cosine columns was not saved as
a standalone source file, so do not claim this file is byte-for-byte original.
"""

import json
import math
import os
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


MODELS = ["GPT-5.6 Sol Ultra", "Fable 5 High", "Grok 4.6 xhigh"]
GROUPS = [
    ("Before (a)", (4, 5, 6)),
    ("Before (b)", (7, 8, 9)),
    ("Before (c)", (10, 11, 12)),
    ("After (a)", (14, 15, 16)),
    ("After (b)", (17, 18, 19)),
    ("After (c)", (20, 21, 22)),
]
TOKEN_PATTERN = re.compile(r"(?u)\b\w\w+\b")


def tokens(text):
    return TOKEN_PATTERN.findall(str(text or "").lower())


def mean(values):
    return sum(values) / len(values) if values else 0.0


def fit_tfidf(corpus):
    counts = [Counter(tokens(document)) for document in corpus]
    document_frequency = Counter()
    for document in counts:
        document_frequency.update(document.keys())

    document_count = len(counts)
    idf = {
        term: math.log((1 + document_count) / (1 + frequency)) + 1
        for term, frequency in document_frequency.items()
    }

    def vectorize(text):
        term_counts = Counter(tokens(text))
        vector = {
            term: frequency * idf.get(term, math.log(1 + document_count) + 1)
            for term, frequency in term_counts.items()
        }
        magnitude = math.sqrt(sum(value * value for value in vector.values()))
        if magnitude:
            vector = {term: value / magnitude for term, value in vector.items()}
        return vector

    return vectorize


def cosine(left, right):
    if len(left) > len(right):
        left, right = right, left
    return sum(value * right.get(term, 0.0) for term, value in left.items())


def main():
    input_path = Path(os.environ["AP_RESEARCH_WORKBOOK"])
    workbook = load_workbook(input_path, read_only=True, data_only=False)
    sheet = workbook[workbook.sheetnames[0]]

    references = {
        str(sheet.cell(row, 42).value or "").strip(): str(sheet.cell(row, 43).value or "")
        for row in range(20, 24)
    }

    candidates = []
    for row in range(2, 18):
        subject = str(sheet.cell(row, 1).value or "").strip()
        for group_name, columns in GROUPS:
            for model, column in zip(MODELS, columns):
                candidates.append({
                    "row": row,
                    "subject": subject,
                    "group": group_name,
                    "model": model,
                    "text": str(sheet.cell(row, column).value or ""),
                })

    # Fit one shared vocabulary/IDF model across every generated response and
    # the four unique teacher-reference documents.
    vectorize = fit_tfidf(
        [item["text"] for item in candidates] + list(references.values())
    )
    reference_vectors = {
        subject: vectorize(text) for subject, text in references.items()
    }

    score_lookup = {}
    for item in candidates:
        score_lookup[(item["row"], item["group"], item["model"])] = cosine(
            vectorize(item["text"]), reference_vectors[item["subject"]]
        )

    headers = [f"Cosine - {group}" for group, _ in GROUPS]
    matrix = []
    for row in range(2, 18):
        output_row = []
        for group, _ in GROUPS:
            output_row.append("\n".join(
                f"{model} - {score_lookup[(row, group, model)]:.4f}"
                for model in MODELS
            ))
        matrix.append(output_row)

    json.dump({"headers": headers, "matrix": matrix}, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main()
